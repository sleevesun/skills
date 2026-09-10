#!/usr/bin/env python3
"""
考勤规则增量校验与确认辅助脚本 (check_rules.py)
用于在运行计算前扫描原始打卡底表，检查是否存在考勤规则库尚未收录的“上午状态|下午状态”组合。
若存在新增组合，主动拦截并输出待确认清单；若用户已确认权重，支持自动更新规则库文件。
"""

import argparse
import json
import sys
from pathlib import Path
from typing import Any, Dict, List, Tuple

try:
    import openpyxl
except ImportError:
    openpyxl = None


def scan_unrecorded_rules(raw_xlsx: Path, rules_json: Path) -> Dict[str, Any]:
    """比对打卡底表与规则库，提取未收录的规则"""
    if openpyxl is None:
        raise RuntimeError("需要 openpyxl 库，请先安装: pip install openpyxl")

    if not raw_xlsx.exists():
        raise FileNotFoundError(f"打卡底表文件不存在: {raw_xlsx}")
    if not rules_json.exists():
        raise FileNotFoundError(f"考勤规则文件不存在: {rules_json}")

    with open(rules_json, "r", encoding="utf-8") as f:
        rules: Dict[str, float] = json.load(f)

    wb = openpyxl.load_workbook(str(raw_xlsx), data_only=True)
    ws = wb["Sheet0"] if "Sheet0" in wb.sheetnames else wb.active

    unknown_rules: Dict[str, Dict[str, Any]] = {}
    matched_pairs: set = set()
    total_rows = 0

    for r in range(2, ws.max_row + 1):
        eid_val = ws.cell(r, 1).value
        if not eid_val or str(eid_val).strip() == "员工编号":
            continue
        total_rows += 1

        eid = str(eid_val).strip()
        name = str(ws.cell(r, 2).value or "").strip()
        dt_str = str(ws.cell(r, 3).value or "")[:10]
        f_am = str(ws.cell(r, 6).value or "").strip()
        g_pm = str(ws.cell(r, 7).value or "").strip()
        p_work = str(ws.cell(r, 16).value or "").strip().upper()

        pair = f"{f_am}|{g_pm}"
        if pair in rules:
            matched_pairs.add(pair)
        else:
            if pair not in unknown_rules:
                unknown_rules[pair] = {
                    "pair": pair,
                    "f_am": f_am,
                    "g_pm": g_pm,
                    "count": 0,
                    "is_workday": p_work != "N",
                    "samples": [],
                }
            unknown_rules[pair]["count"] += 1
            if len(unknown_rules[pair]["samples"]) < 3:
                unknown_rules[pair]["samples"].append({
                    "eid": eid,
                    "name": name,
                    "date": dt_str,
                    "workday": p_work != "N",
                })

    wb.close()
    return {
        "ok": len(unknown_rules) == 0,
        "total_rows": total_rows,
        "matched_rules_count": len(matched_pairs),
        "unknown_count": len(unknown_rules),
        "unknown_rules": unknown_rules,
    }


def add_rule_to_json(rules_json: Path, pair: str, weight: float) -> None:
    """向规则库追加已确认的新考勤规则"""
    with open(rules_json, "r", encoding="utf-8") as f:
        rules = json.load(f)

    rules[pair] = float(weight)

    with open(rules_json, "w", encoding="utf-8") as f:
        json.dump(rules, f, ensure_ascii=False, indent=2)

    print(f"✓ 规则库已更新: 追加规则「{pair}」= {weight}")


def format_unknown_rules_prompt(result: Dict[str, Any]) -> str:
    """生成面向用户的结构化确认提示语"""
    if result["ok"]:
        return (
            f"✅ **考勤规则核验通过**：\n"
            f"打卡底表共扫描 {result['total_rows']} 条打卡明细，涵盖的所有考勤状态组合（共 {result['matched_rules_count']} 种）"
            f"均已 100% 收录在规则库中，无任何新增或未定义规则。"
        )

    lines = []
    lines.append("⚠️ **考勤规则增量校验拦截：发现未收录的新考勤状态组合！**")
    lines.append(f"在原始打卡底表中扫描到 **{result['unknown_count']} 种** 尚未收录在 `考勤规则.json` 中的新组合：\n")

    for idx, (pair, info) in enumerate(result["unknown_rules"].items(), 1):
        lines.append(f"{idx}. 新状态组合：**`{pair}`**")
        lines.append(f"   - 累计出现频次：{info['count']} 次")
        samples_str = "、".join([f"{s['name']}({s['eid']}, {s['date']})" for s in info['samples']])
        lines.append(f"   - 涉及样本员工：{samples_str}")
        lines.append(f"   - 是否工作日：{'工作日' if info['is_workday'] else '非工作日/周末'}")
        lines.append("")

    lines.append("👉 **需要业务确认**：")
    lines.append("请确认上述新组合的出勤折算权重：")
    lines.append("  • `1.0`：计为全天出勤")
    lines.append("  • `0.5`：计为半天出勤（如已请假、已冲销、迟到等配合正常打卡）")
    lines.append("  • `0.0`：不计出勤天数（如请假转出勤、全天缺勤旷工等）")
    lines.append("\n请回复确认各组合对应的权重值，系统将在录入规则库后继续执行后续核算。")
    return "\n".join(lines)


def main():
    parser = argparse.ArgumentParser(description="校验打卡底表中的考勤状态组合是否收录在规则库中")
    parser.add_argument("--raw", required=True, help="原始考勤打卡底表 (.xlsx)")
    parser.add_argument("--rules", required=True, help="考勤规则库 (.json)")
    parser.add_argument("--add", help="已确认追加的新规则组合，如 '请假转出勤|请假转出勤'")
    parser.add_argument("--weight", type=float, help="配合 --add 传入的权重值，如 0.0, 0.5, 1.0")
    parser.add_argument("--json", action="store_true", help="以 JSON 格式输出")
    args = parser.parse_args()

    raw_path = Path(args.raw).resolve()
    rules_path = Path(args.rules).resolve()

    if args.add and args.weight is not None:
        add_rule_to_json(rules_path, args.add, args.weight)
        return

    res = scan_unrecorded_rules(raw_path, rules_path)
    if args.json:
        print(json.dumps(res, ensure_ascii=False, indent=2))
    else:
        print(format_unknown_rules_prompt(res))


if __name__ == "__main__":
    main()
